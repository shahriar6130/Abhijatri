"""Helpers for building the Abhijatri Word documents with python-docx (A4, consistent styles, tables, figures)."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

BLUE = RGBColor(0x00, 0x55, 0x8C)
GREY = RGBColor(0x55, 0x55, 0x55)
HEAD_FILL = "DCEBF7"
NOTE_FILL = "FFF4D6"
CODE_FILL = "F1F1F1"
GOOD_FILL = "E3F3EA"
BAD_FILL = "F9DFD8"


PPR_AFTER_SHD = ("w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE",
                 "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing",
                 "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
                 "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")
PPR_AFTER_BDR = ("w:shd",) + PPR_AFTER_SHD
TBLPR_AFTER_BORDERS = ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription", "w:tblPrChange")
SETTINGS_AFTER_UPDATE = ("w:hdrShapeDefaults", "w:footnotePr", "w:endnotePr", "w:compat", "w:docVars", "w:rsids", "m:mathPr",
                         "w:attachedSchema", "w:themeFontLang", "w:clrSchemeMapping", "w:doNotIncludeSubdocsInStats",
                         "w:doNotAutoCompressPictures", "w:forceUpgrade", "w:captions", "w:readModeInkLockDown", "w:smartTagType",
                         "sl:schemaLibrary", "w:shapeDefaults", "w:doNotEmbedSmartTags", "w:decimalSymbol", "w:listSeparator")


def _insert_ordered(parent, element, successors) -> None:
    """Insert element before the first existing child that must follow it in the schema; else append."""
    parent.insert_element_before(element, *successors)


class Doc:
    """A document with running figure/table numbers and helpers; portrait A4 by default."""

    def __init__(self, title: str, subtitle: str, meta: list[tuple[str, str]], footer_text: str) -> None:
        self.d = Document()
        self.fig = 0
        self.tab = 0
        self.footer_text = footer_text
        self._styles()
        self._page(portrait=True)
        self._footer(self.d.sections[0])
        self._title_page(title, subtitle, meta)

    # ---------------------------------------------------------------- setup
    def _styles(self) -> None:
        st = self.d.styles
        normal = st["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(10.5)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.12
        for name, size, before, after in (("Heading 1", 18, 18, 8), ("Heading 2", 14, 14, 6), ("Heading 3", 12, 10, 4)):
            h = st[name]
            h.font.name = "Calibri"
            h.font.size = Pt(size)
            h.font.bold = True
            h.font.color.rgb = BLUE
            h.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
            h.paragraph_format.space_before = Pt(before)
            h.paragraph_format.space_after = Pt(after)
            h.paragraph_format.keep_with_next = True
        for name in ("List Bullet", "List Bullet 2"):
            st[name].font.name = "Calibri"
            st[name].font.size = Pt(10.5)
            st[name].paragraph_format.space_after = Pt(3)

    def _page(self, portrait: bool) -> None:
        s = self.d.sections[-1]
        s.page_width, s.page_height = (Cm(21.0), Cm(29.7)) if portrait else (Cm(29.7), Cm(21.0))
        s.orientation = WD_ORIENT.PORTRAIT if portrait else WD_ORIENT.LANDSCAPE
        s.left_margin = s.right_margin = Cm(2.2)
        s.top_margin = Cm(2.0)
        s.bottom_margin = Cm(2.0)

    def _footer(self, section) -> None:
        section.footer.is_linked_to_previous = False
        p = section.footer.paragraphs[0]
        p.text = ""
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(self.footer_text + "   |   Page ")
        r.font.size = Pt(8.5)
        r.font.color.rgb = GREY
        self._field(p, "PAGE", size=8.5)

    def _field(self, paragraph, instr: str, size: float = 10.0, placeholder: str = "") -> None:
        run = paragraph.add_run()
        run.font.size = Pt(size)
        run.font.color.rgb = GREY if size < 9 else None
        for kind, text in (("begin", None), (None, instr), ("separate", None), ("text", placeholder), ("end", None)):
            if kind in ("begin", "separate", "end"):
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), kind)
                run._r.append(el)
            elif kind == "text":
                t = OxmlElement("w:t")
                t.text = text
                run._r.append(t)
            else:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = f" {text} "
                run._r.append(el)

    def _title_page(self, title: str, subtitle: str, meta: list[tuple[str, str]]) -> None:
        for _ in range(5):
            self.d.add_paragraph()
        p = self.d.add_paragraph()
        r = p.add_run("Abhijatri")
        r.font.size = Pt(34)
        r.font.bold = True
        r.font.color.rgb = BLUE
        p = self.d.add_paragraph()
        r = p.add_run(title)
        r.font.size = Pt(24)
        r.font.bold = True
        p = self.d.add_paragraph()
        r = p.add_run(subtitle)
        r.font.size = Pt(13)
        r.font.color.rgb = GREY
        self.d.add_paragraph()
        self.table(["Item", "Detail"], [[a, b] for a, b in meta], [4.0, 12.5], caption=None, size=10)
        self.page_break()

    # ---------------------------------------------------------------- text
    def h1(self, text: str) -> None:
        self.d.add_heading(text, level=1)

    def h2(self, text: str) -> None:
        self.d.add_heading(text, level=2)

    def h3(self, text: str) -> None:
        self.d.add_heading(text, level=3)

    def p(self, text: str = "", bold: bool = False, italic: bool = False, size: float | None = None, color: RGBColor | None = None, align=None, after: float | None = None):
        para = self.d.add_paragraph()
        self._runs(para, text, bold, italic, size, color)
        if align:
            para.alignment = align
        if after is not None:
            para.paragraph_format.space_after = Pt(after)
        return para

    def _runs(self, para, text: str, bold=False, italic=False, size=None, color=None) -> None:
        """Supports **bold** and `code` spans inside text."""
        import re

        for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
            if not part:
                continue
            if part.startswith("**") and part.endswith("**"):
                r = para.add_run(part[2:-2])
                r.bold = True
            elif part.startswith("`") and part.endswith("`"):
                r = para.add_run(part[1:-1])
                r.font.name = "Consolas"
                r.font.size = Pt((size or 10.5) - 1)
                r.element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
                if bold:
                    r.bold = True
                continue
            else:
                r = para.add_run(part)
                r.bold = bold or None
            r.italic = italic or None
            if size:
                r.font.size = Pt(size)
            if color is not None:
                r.font.color.rgb = color

    def bullets(self, items: list[str], level: int = 1) -> None:
        for item in items:
            para = self.d.add_paragraph(style="List Bullet" if level == 1 else "List Bullet 2")
            self._runs(para, item)

    def steps(self, items: list[str]) -> None:
        """Manually numbered steps (avoids Word's list numbering carrying over between lists)."""
        for i, item in enumerate(items, 1):
            para = self.d.add_paragraph()
            para.paragraph_format.left_indent = Cm(0.9)
            para.paragraph_format.first_line_indent = Cm(-0.9)
            para.paragraph_format.space_after = Pt(3)
            r = para.add_run(f"{i}.\t")
            r.bold = True
            self._runs(para, item)

    def code(self, text: str) -> None:
        for line in text.strip("\n").split("\n"):
            para = self.d.add_paragraph()
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.left_indent = Cm(0.4)
            r = para.add_run(line if line else " ")
            r.font.name = "Consolas"
            r.font.size = Pt(8.8)
            r.element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            self._shade_paragraph(para, CODE_FILL)
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)

    def note(self, text: str, fill: str = NOTE_FILL) -> None:
        para = self.d.add_paragraph()
        para.paragraph_format.left_indent = Cm(0.3)
        para.paragraph_format.right_indent = Cm(0.3)
        self._runs(para, text)
        self._shade_paragraph(para, fill)

    def _shade_paragraph(self, para, fill: str) -> None:
        pPr = para._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        _insert_ordered(pPr, shd, PPR_AFTER_SHD)

    def page_break(self) -> None:
        self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---------------------------------------------------------------- contents
    def contents(self) -> None:
        self.h1("Contents")
        para = self.d.add_paragraph()
        self._field(para, 'TOC \\o "1-2" \\h \\z \\u', size=10.5,
                    placeholder="Contents appear here. If this line is still visible, right-click it and choose Update Field (or press F9).")
        settings = self.d.settings.element
        flag = OxmlElement("w:updateFields")
        flag.set(qn("w:val"), "true")
        _insert_ordered(settings, flag, SETTINGS_AFTER_UPDATE)
        for zoom in settings.findall(qn("w:zoom")):
            if zoom.get(qn("w:percent")) is None:
                zoom.set(qn("w:percent"), "100")
        self.page_break()

    # ---------------------------------------------------------------- tables
    def table(self, header: list[str], rows: list[list[str]], widths_cm: list[float], caption: str | None = None, size: float = 9.0,
              fills: dict[tuple[int, int], str] | None = None, bold_first_col: bool = False) -> None:
        if caption:
            self.tab += 1
            cp = self.d.add_paragraph()
            r = cp.add_run(f"Table {self.tab}. {caption}")
            r.bold = True
            r.font.size = Pt(9.5)
            r.font.color.rgb = BLUE
            cp.paragraph_format.keep_with_next = True
            cp.paragraph_format.space_after = Pt(3)
        t = self.d.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        for i, text in enumerate(header):
            c = t.rows[0].cells[i]
            self._cell(c, text, widths_cm[i], size, bold=True, fill=HEAD_FILL)
        trPr = t.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader")
        th.set(qn("w:val"), "true")
        trPr.append(th)
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            for ci, text in enumerate(row):
                fill = (fills or {}).get((ri, ci))
                self._cell(cells[ci], str(text), widths_cm[ci], size, bold=bold_first_col and ci == 0, fill=fill)
            trPr = t.rows[-1]._tr.get_or_add_trPr()
            cs = OxmlElement("w:cantSplit")
            cs.set(qn("w:val"), "true")
            trPr.append(cs)
        self.d.add_paragraph().paragraph_format.space_after = Pt(4)

    def _table_borders(self, table) -> None:
        tblPr = table._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{edge}")
            for k, v in (("val", "single"), ("sz", "4"), ("space", "0"), ("color", "9AA5B1")):
                el.set(qn(f"w:{k}"), v)
            borders.append(el)
        _insert_ordered(tblPr, borders, TBLPR_AFTER_BORDERS)

    def _cell(self, cell, text: str, width_cm: float, size: float, bold: bool = False, fill: str | None = None) -> None:
        cell.width = Cm(width_cm)
        tcPr = cell._tc.get_or_add_tcPr()
        mar = OxmlElement("w:tcMar")
        for side, width in (("top", "50"), ("left", "90"), ("bottom", "50"), ("right", "90")):
            el = OxmlElement(f"w:{side}")
            el.set(qn("w:w"), width)
            el.set(qn("w:type"), "dxa")
            mar.append(el)
        if fill:
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), fill)
            tcPr.append(shd)
        tcPr.append(mar)
        lines = text.split("\n")
        para = cell.paragraphs[0]
        for i, line in enumerate(lines):
            if i:
                para = cell.add_paragraph()
            para.paragraph_format.space_after = Pt(1)
            para.paragraph_format.line_spacing = 1.0
            self._runs(para, line, bold=bold, size=size)

    # ---------------------------------------------------------------- figures and sections
    def landscape(self) -> None:
        s = self.d.add_section(WD_SECTION.NEW_PAGE)
        self._page(portrait=False)
        self._footer(s)

    def portrait(self) -> None:
        s = self.d.add_section(WD_SECTION.NEW_PAGE)
        self._page(portrait=True)
        self._footer(s)

    def figure(self, path: str | Path, caption: str, width_cm: float, max_height_cm: float | None = None) -> None:
        path = Path(path)
        with Image.open(path) as im:
            w, h = im.size
        height = width_cm * h / w
        if max_height_cm and height > max_height_cm:
            width_cm = width_cm * max_height_cm / height
        self.fig += 1
        para = self.d.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.keep_with_next = True
        para.paragraph_format.space_after = Pt(2)
        para.add_run().add_picture(str(path), width=Cm(width_cm))
        cp = self.d.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cp.add_run(f"Figure {self.fig}. {caption}")
        r.italic = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = GREY

    def save(self, path: str | Path) -> None:
        self.d.save(str(path))
